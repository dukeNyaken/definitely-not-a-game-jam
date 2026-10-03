"""Строит таблицу «ситуация → диалог» (docs/Сюжет_диалоги.xlsx) из выгрузки сюжета.

    godot --headless --path . -s tools/export_story.gd -- out=story.json
    python tools/story_table.py story.json docs/Сюжет_диалоги.xlsx [прежняя_таблица.xlsx]

Третий аргумент — прежняя версия таблицы: изменившиеся реплики получают в колонке «Было раньше»
старый текст, новые строки помечаются. Нужен openpyxl.
"""
import json
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
prev_path = sys.argv[3] if len(sys.argv) > 3 else None
lines = json.load(open(src, encoding="utf-8"))

ITEMS = [("sword", "Меч", "friend"), ("amulet", "Амулет (оберег)", "beloved"), ("boots", "Сапоги", "refugee"),
         ("gloves", "Перчатки с перстнем отца", "captain"), ("shield", "Щит", "widow"), ("armor", "Доспех", "smith"),
         ("helmet", "Шлем", "novice")]
WHO = ["hero", "father", "brother", "beloved", "mother", "faithful", "friend", "refugee", "captain", "widow", "smith", "novice"]
# Метки, переименованные с прежней версии таблицы: новая -> старая.
RENAMED = {"gates.they_came": "gates.waited"}
NEW_MARK = "— новая строка —"


def name(who):
    if who in ("thought", ""):
        return {"thought": "Солдат (мысль)", "": "—"}[who]
    return lines["speaker.%s.name" % who]["text"]


rows = []  # (метка, сцена, ситуация, условие, кто, вид)


def add(label, scene, situation, condition="", who=None, kind="реплика"):
    if label not in lines:
        raise SystemExit("нет метки " + label)
    if who is None:
        who = lines[label]["speaker"]
    rows.append((label, scene, situation, condition, name(who) if kind != "летопись" else "летопись", kind))


def chapter(key, scene, situation):
    add("chapter.%s.title" % key, scene, situation, who="", kind="заставка")
    add("chapter.%s.sub" % key, scene, "Подзаголовок той же заставки (мелко, под заголовком).", who="", kind="заставка")


S0 = "0. Персонажи"
for who in WHO:
    add("speaker.%s.name" % who, S0, "Имя: подпись над репликами и крупная строка таблички над головой.", who="", kind="имя")
    add("speaker.%s.role" % who, S0, "Роль: мелкая строка таблички под именем, когда персонаж впервые появляется в кадре.", who="", kind="роль")

S1 = "1. Пролог"
P = "prologue."
chapter("prologue", S1, "Начало забега. Заставка главы на чёрном экране: золотой заголовок между двумя линиями.")
add(P + "two_sons", S1, "Чёрный экран после заставки.", kind="летопись")
add(P + "young_hero_name", S1, "Флешбэк на «старой плёнке» (сепия, зерно, мерцание): солнечный двор крепости, стойка с деревянными мечами, чучело. "
    "Ярл смотрит, как сыновья бьются. Табличка над младшим (имя).", who="", kind="табличка")
add(P + "young_hero_role", S1, "Та же табличка над младшим (роль).", who="", kind="табличка")
add(P + "young_brother_role", S1, "Табличка над старшим: имя берётся из speaker.brother.name, здесь — роль.", who="", kind="табличка")
add(P + "spar_rise", S1, "Три обмена ударами (щепки, тряска кадра). Последний удар младшего — замедленно, вспышка; старший падает в пыль. "
    "Младший бросает меч и протягивает брату руку.", who="hero")
add(P + "spar_refuse", S1, "Сигвард отбивает протянутую руку и встаёт сам.", who="brother")
add(P + "father_spar", S1, "Камера на ярле. Его слова — ключ к финалу: старший смотрит на вещи, младший — на людей.", who="father")
add(P + "ring_intro", S1, "Затемнение. Годы спустя.", kind="летопись")
add(P + "father_ring", S1, "Покои ярла при свече: сыновья стоят перед сгорбленным отцом. Крупно — светящийся перстень в его протянутой руке.", who="father")
add(P + "father_ring_2", S1, "Ярл поворачивается к младшему. После реплики перстень перелетает к нему: золотой шлейф, столб света.", who="father")
add(P + "brother_envy", S1, "Крупно Сигвард, медленный наезд. После реплики он уходит в темноту.", who="brother")
add(P + "father_death", S1, "Ярл оседает и растворяется; свеча гаснет вместе с ним. Экран темнеет.", kind="летопись")
add(P + "tyrant", S1, "Чёрный экран, конец флешбэка.", kind="летопись")
add(P + "beloved_1", S1, "Настоящее: камера обходит Солдата, семь вещей по очереди вспыхивают цветами своих сущностей, последним — перстень. "
    "Подходит Сольвейг (табличка).", who="beloved")
add(P + "hero_1", S1, "Ответ Сольвейг.", who="hero")
add(P + "beloved_2", S1, "Двойной смысл: «у ворот» окажется воротами дворца.", who="beloved")
add(P + "faithful_1", S1, "Из-за спины выходит Ильва с фонарём (табличка) и протягивает узелок. Сольвейг ещё здесь.", who="faithful")
add(P + "beloved_3", S1, "Сольвейг — Ильве, свысока. После реплики уходит в темноту.", who="beloved")
add(P + "hero_2", S1, "Солдат берёт хлеб у Ильвы.", who="hero")
add(P + "learn", S1, "Ильва не уходит — встаёт в паре шагов за его спиной. Кадр темнеет. После — рог и баннер «Этап 1».", kind="летопись")

S1B = "1б. Первый алтарь (после первого этапа)"
chapter("rule", S1B, "Первый этап пройден, посреди арены поднимается алтарь. Солдат подходит к нему; компактная заставка в верхней части кадра.")
add("speaker.chronicle.name", S1B, "Чьим голосом идёт предание: подпись над его репликами.", who="", kind="имя")
add("rule.law", S1B, "Свет алтаря поднимается столбом. В нём встают две золотые тени прошлого: воин протягивает меч тому, кому он нужнее, — "
    "и меч перелетает из рук в руки. Правило мира (формулировка автора — менять только осознанно).",
    "один раз за забег — у первого алтаря", who="", kind="предание")
add("rule.light", S1B, "Сила меча светом возвращается к отдавшему. От его ног загорается дорога из света, он уходит по ней; тени тают. "
    "Вторая половина правила.", "один раз за забег — у первого алтаря", who="", kind="предание")

S2 = "2. Дар (после жертвы на алтаре)"
add("chapter.gift.title", S2, "Компактная заставка в верхней части кадра, пока получатель идёт к Солдату. "
    "{n} — порядковое слово (ниже); скобки не удалять.", who="", kind="заставка")
add("chapter.gift.sub", S2, "Подзаголовок: {item} — вещь с заглавной буквы, {name} — имя получателя.", who="", kind="заставка")
for n in range(1, 7):
    add("ordinal.%d" % n, S2, "Порядковое слово для {n} в заставке дара №%d («Дар первый»)." % n, who="", kind="словоформа")
for item, title, who in ITEMS:
    add("item.%s" % item, S2, "«%s» в сюжете: для {item} в заставке и для {recipient} в подписи о силе "
        "(именительный падеж, со строчной)." % title, who="", kind="словоформа")
add("caption.light", S2, "Столб света и поток силы из вещи в руках получателя уходят в соседнюю вещь Солдата. "
    "Подстановки {victim} {recipient} {property} — фигурные скобки не удалять.", kind="мысль")
for item, title, who in ITEMS:
    add("gen.%s" % item, S2, "Родительный падеж «%s» для {victim} в подписи выше («сила меча»)." % title, who="", kind="словоформа")
for item, title, who in ITEMS:
    scene = "2. Дар: %s → %s" % (title, lines["speaker.%s.name" % who]["text"])
    cond = "если отдана вещь «%s»" % title
    add("gift.%s.plea" % item, scene, "Получатель выходит из темноты к Солдату у алтаря (табличка с именем) и просит. Камера близко, медленный облёт.", cond)
    add("gift.%s.reply" % item, scene, "Ответ Солдата. Затем вещь перелетает к получателю со шлейфом цвета своей сущности.", cond)
    add("gift.%s.oath" % item, scene, "Получатель клянётся в низком поклоне, прижав вещь к груди.", cond)
    if "gift.%s.extra" % item in lines:
        add("gift.%s.extra" % item, scene, "Мысль после подписи о силе (только для перчаток с перстнем).", cond, kind="мысль")
    add("gift.%s.secret" % item, scene, "Крупный план вещи в чужих руках. Мысль: что Солдат знает о её слабом месте — закладка для финала.", cond, kind="мысль")
IVA_ACTS = {1: "Ильва отдаёт Солдату свою флягу.", 3: "Ильва отдаёт Солдату свой хлеб."}
for n in range(1, 7):
    add("iva.%d" % n, S2, "Конец сцены дара №%d (какая бы вещь ни была отдана). Камера на Ильве и Солдате. %s" % (n, IVA_ACTS.get(n, "")),
        "дар №%d" % n)

S3 = "3. У очага (после второго дара)"
H = "hearth."
chapter("hearth", S3, "Между этапами, после дара №2 и перед голосом из дворца. Компактная заставка в верхней части кадра.")
add(H + "cant_sleep", S3, "Дом Сольвейг, ночь: огонь в очаге, свеча на столе, окно в холодную ночь. Мать прядёт у огня, "
    "Сольвейг стоит у окна спиной к дому. Мать оставляет прялку и оборачивается (табличка).", who="mother")
add(H + "love", S3, "Сольвейг оборачивается от окна (табличка). Крупно, медленный облёт.", who="beloved")
add(H + "then_why", S3, "Общий план: мать и дочь.", who="mother")
add(H + "too_good", S3, "Сольвейг идёт от окна к столу, к свече.", who="beloved")
add(H + "gave", S3, "{items} — вещи, которые Солдат уже отдал чужим (без оберега): «меч», «меч и сапоги». Скобки не удалять.",
    "если чужим отдана хоть одна вещь", who="beloved")
add(H + "amulet", S3, "Крупно — оберег Солдата в её руках; он вспыхивает своим цветом.", "если оберег уже отдан Сольвейг", who="beloved")
add(H + "squander", S3, "Она оборачивается к сундуку с приданым.", who="beloved")
add(H + "kindness", S3, "Мать подходит к дочери.", who="mother")
add(H + "not_feed", S3, "Сольвейг отворачивается от матери — лицом к зрителю. Кадр холодеет, огонь в очаге садится.", who="beloved")
add(H + "not_poverty", S3, "Крупно Сольвейг, медленный наезд.", who="beloved")
add(H + "alone", S3, "Тот же кадр: она опускает голову. Главная реплика сцены.", who="beloved")
add(H + "tell_him", S3, "Мать протягивает к ней руку.", who="mother")
add(H + "wont_hear", S3, "Сольвейг — матери. После реплики возвращается к окну; свеча на столе гаснет, экран темнеет.", who="beloved")

S3B = "4. Тронный зал: Сигвард и Сольвейг (после третьего дара)"
T = "temptation."
chapter("temptation", S3B, "Между этапами, после дара №3 и перед голосом из дворца. Компактная заставка в верхней части кадра.")
add(T + "summon", S3B, "Тронный зал в багровом свете жаровен. Сольвейг идёт по ковру к трону; Сигвард спускается ей навстречу "
    "(таблички).", who="brother")
add(T + "speak", S3B, "Крупно Сольвейг.", who="beloved")
add(T + "gives_away", S3B, "Они стоят лицом к лицу поперёк ковра; камера медленно обходит их.", who="brother")
add(T + "needier", S3B, "Ответ Сольвейг.", who="beloved")
add(T + "and_you", S3B, "Сигвард обходит её и встаёт за плечом; она отворачивается и опускает голову. После реплики — пауза: "
    "кадр на миг холодеет, ей нечего ответить.", who="brother")
add(T + "leave_him", S3B, "В руке Сигварда кошель с золотом. После реплики он бросает кошель к её ногам: монеты, блик.", who="brother")
add(T + "no", S3B, "Сольвейг выпрямляется и оборачивается к нему. Отказ.", who="beloved")
add(T + "love_winter", S3B, "Крупно Сигвард.", who="brother")
add(T + "dont_call", S3B, "Она уходит. У кошеля замедляет шаг и смотрит на золото (первый раз). У края ковра оборачивается "
    "к трону (второй раз) — и говорит.", who="beloved")
add(T + "looked_back", S3B, "Сольвейг уходит в темноту. Сигвард один над кошелем; жаровни разгораются, медленный наезд.", who="brother")

S3 = "5. Голос из дворца"
add("caption.interlude_header", S3, "Подпись места в верхней части кадра. {name} — имя Сигварда, скобки не удалять.", who="", kind="заставка")
PALACE = {
    1: "У окна: смотрит на зарево над городом, на середине реплики оборачивается.",
    2: "Идёт по ковру от трона прямо на камеру, руки за спиной.",
    3: "У жаровни, в руке детский деревянный меч. После реплики бросает его в огонь — пламя взвивается.",
    4: "Читает письмо Сольвейг с багровой печатью (после разговора в тронном зале она пишет ему сама). Дочитав, роняет его.",
    5: "Смотрит на свою пустую руку, на которой нет перстня.",
    6: "Встаёт перед троном, раскидывает руки — обе жаровни вспыхивают.",
}
for n in range(1, 7):
    cond = "дар №%d" % n
    if n == 5:
        cond += ", перчатки ещё НЕ отданы"
    add("brother.%d" % n, S3, "Между этапами: тронный зал в багровом свете жаровен. " + PALACE[n], cond)
    if "brother.%d.ring" % n in lines:
        add("brother.%d.ring" % n, S3, "Тот же кадр: " + PALACE[n], "дар №%d, перчатки с перстнем уже отданы Бьёрну" % n)
    if "brother.%d.gum" % n in lines:
        add("brother.%d.gum" % n, S3, "Шутка после основной реплики. У модели Сигварда за поднятой рукой от бедра тянется нить; здесь она — розовая жвачка "
            "из кармана. Камера подходит ближе, он говорит о ней сам, опускает руки — и жвачка лопается.", "дар №%d" % n)

S4 = "6. У ворот дворца"
G = "gates."
chapter("gates", S4, "Последний этап. Заставка главы на чёрном экране.")
add(G + "they_came", S4, "Общий план с облётом: закрытые ворота в багровом зареве, на постаментах светятся отданные вещи, "
    "у постаментов — получатели спиной к Солдату.", who="thought", kind="мысль")
for item, title, who in ITEMS:
    if "gift.%s.gate" % item not in lines:
        continue
    add("gift.%s.gate" % item, S4, "Камера наезжает на получателя (табличка). Он оборачивается к Солдату, говорит и снова отворачивается. "
        "Звучат не больше трёх таких реплик — у первых по порядку жертв.", "если отдана вещь «%s»" % title)
add(G + "one_brother", S4, "Ворота открылись (грохот, пыль, багровый свет изнутри). Сигвард вышел и встал напротив Солдата.",
    "если у ворот говорил Торстейн (он зовёт Солдата «брат»)", who="brother")
add(G + "hello", S4, "Сигвард напротив Солдата (табличка).", who="brother")
add(G + "envy_1", S4, "Монолог зависти: камера медленно наезжает на Сигварда и обходит его.", who="brother")
add(G + "envy_2", S4, "Монолог, продолжение.", who="brother")
add(G + "envy_3", S4, "Монолог, конец. Вещи ещё на постаментах — он говорит о том, что сейчас будет.", who="brother")
add(G + "not_things", S4, "Крупно Солдат.", who="hero")
add(G + "we_will_see", S4, "Крупно Сигвард.", who="brother")
add(G + "sol_offer", S4, "Сольвейг вышла из ворот и встала рядом с Сигвардом (табличка). После реплики поднимает руки — "
    "вещи слетаются с постаментов в Сигварда, получатели отворачиваются к нему.", who="beloved")
add(G + "ring_at_last", S4, "Вещи собрались на Сигварде; золотой блик на его руке.", "если отданы перчатки", who="brother")
add(G + "confess_given", S4, "Признание Сольвейг, крупно, медленный облёт: она любила — и выбрала дворец, а не нищету.", "если оберег отдан", who="beloved")
add(G + "confess_kept", S4, "Признание Сольвейг, крупно, медленный облёт: она любила — и выбрала дворец, а не нищету.", "если оберег НЕ отдан", who="beloved")
add(G + "hurt", S4, "Крупно Солдат; кадр на миг холодеет и обесцвечивается. Отсылка к правилу мира.", who="thought", kind="мысль")
add(G + "unbreakable", S4, "Превращение: замедление, столб света, белая вспышка — на месте Сигварда рыцарь-босс.", who="brother")
add(G + "here", S4, "Предатели отходят к воротам. Ильва с фонарём выходит из-за спины Солдата.", who="faithful")
add(G + "i_know", S4, "Ответ Ильве.", who="hero")
add(G + "know_things", S4, "Ильва отходит к краю арены. Последняя мысль перед боем.", who="thought", kind="мысль")
add(G + "banner_sub", S4, "Баннер в начале боя под крупным словом «Тиран».", who="", kind="баннер")

S5 = "7. Бой с Тираном"
add("bark.brother.1", S5, "Субтитр в бою (игра не останавливается): переход во 2-ю фазу, Тиран сбрасывает вещи.", kind="субтитр")
add("bark.brother.2", S5, "Субтитр в бою: переход в 3-ю фазу.", kind="субтитр")
add("bark.iva", S5, "Субтитр в бою: переход в 3-ю фазу, после крика Сигварда.", kind="субтитр")
for item, title, who in ITEMS:
    add("gift.%s.bark" % item, S5, "Субтитр-мысль: Тиран только что сбросил эту вещь.", "если отдана вещь «%s»" % title, kind="субтитр")

S6 = "8. Финал"
F = "finale."
add(F + "how", S6, "Рыцарь рассыпается в замедлении (вспышка, кольцо, пыль); на его месте на коленях Сигвард — просто человек. Солдат подходит.", who="brother")
for item, title, who in ITEMS:
    add("gift.%s.payoff" % item, S6, "Монтаж: отданные вещи лежат вокруг Сигварда. Камера по очереди подходит к каждой (до 4 вещей), "
        "вещь вспыхивает цветом своей сущности.", "если отдана вещь «%s»" % title, kind="мысль")
add(F + "never_helped", S6, "Солдат над братом. Главная реплика истории.", who="hero")
add(F + "father_knew", S6, "Крупно Сигвард.", who="brother")
add(F + "captain_guard", S6, "Перстень лежит на камнях за спиной Сигварда. Солдат делает шаг к брату — камера уходит к воротам: Бьёрн (табличка) "
    "решает, что Солдат идёт добивать, и поднимает щит.", "если отданы перчатки", who="captain")
add(F + "captain", S6, "Бьёрн вбегает между братьями; Солдат отшатывается на шаг. После мысли он заходит слева: искры, щит падает на камни.",
    "если отданы перчатки", who="thought", kind="мысль")
add(F + "captain_calm", S6, "Солдат — Бьёрну, над упавшим щитом. Бьёрн отходит в сторону и опускает голову; Солдат идёт за перстнем.",
    "если отданы перчатки", who="hero")
add(F + "ring_gift", S6, "Солдат протягивает брату перстень отца (поднял с камней или снял со своей перчатки). "
    "После реплики перстень перелетает к Сигварду: золотой столб света.", who="hero")
add(F + "why", S6, "Крупно Сигвард с перстнем в руке.", who="brother")
add(F + "cant_otherwise", S6, "Крупно Солдат. После реплики Сигвард сникает.", who="hero")
add(F + "sorry", S6, "Общий план: те, кого он спас, стоят у ворот и оборачиваются к Солдату. Один делает шаг вперёд (табличка). "
    "Тем временем Сольвейг выходит из толпы и идёт в обход братьев.", "говорит Эйвинд; если его нет — первый из стоящих у ворот", who="refugee")
add(F + "not_for_thanks", S6, "Ответ Солдата.", who="hero")
add(F + "my_soldier", S6, "Сольвейг стоит сбоку от его пути и протягивает руку. Он молча проходит мимо, на миг остановившись рядом; она опускает голову.", who="beloved")
add(F + "lets_go", S6, "Ильва идёт от края арены ему навстречу (табличка); кадр теплеет, как на рассвете. Спасённые расходятся от ворот в темноту.", who="faithful")
add(F + "lets_go_2", S6, "Ответ Солдата. Они уходят вместе, Сольвейг — своей дорогой. Последний кадр — Сигвард один среди вещей.", who="hero")
chapter("epilogue", S6, "Затемнение. Заставка перед финальной летописью.")
add(F + "end_1", S6, "Чёрный экран.", kind="летопись")
add(F + "end_2", S6, "Летопись, продолжение.", kind="летопись")
add(F + "end_3", S6, "Летопись, конец.", kind="летопись")
add(F + "epilogue", S6, "Карточка победы, под строкой «Тиран повержен».", who="", kind="подпись")

missing = sorted(set(lines) - {r[0] for r in rows})
if missing:
    raise SystemExit("метки без строки в таблице: %s" % missing)

# --- Прежняя версия: что изменилось ---------------------------------------
old_text = None
if prev_path:
    old_text = {}
    prev = load_workbook(prev_path)["Диалоги"]
    head = [c.value for c in prev[1]]
    li, ti = head.index("Метка"), head.index("Текст сейчас")
    for r in prev.iter_rows(min_row=2, values_only=True):
        if r[li]:
            old_text[r[li]] = r[ti] or ""


def before(label):
    """Прежний текст, если он изменился; пометка для новой строки; пусто, если всё по-старому."""
    if old_text is None:
        return None
    key = RENAMED.get(label, label)
    if key not in old_text:
        return NEW_MARK
    return old_text[key] if old_text[key] != lines[label]["text"] else None


# --- Книга ------------------------------------------------------------------
FONT = "Arial"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="3A2A2A")
edit_fill = PatternFill("solid", fgColor="FFF2CC")
changed_fill = PatternFill("solid", fgColor="E2EFDA")
scene_fills = ["F2F2F2", "FBEFEF", "EEF3FB", "FDF1E7", "F6EEF3", "F3EEFA", "EEF7EE", "FBF5EA", "F6F6F6"]
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
info = wb.active
info.title = "Как заполнять"
info.sheet_view.showGridLines = False
info["A1"] = "Сюжет «Только самое нужное» — таблица реплик"
info["A1"].font = Font(name=FONT, size=16, bold=True)
notes = [
    "Лист «Диалоги» — все тексты сцен по порядку: что на экране → кто говорит → текст.",
    "Чтобы изменить текст, пишите новый вариант в жёлтую колонку «Новый текст». Пустая клетка — текст остаётся как есть.",
    "«Комментарий» (тоже жёлтая) — для пожеланий к сцене: камера, паузы, что непонятно и т. п.",
    "Колонку «Метка» не меняйте: по ней правки возвращаются в игру (scripts/story/story.gd).",
    "Чтобы убрать реплику, напишите в «Новый текст» слово УДАЛИТЬ.",
    "Фигурные скобки {victim}, {recipient}, {property}, {name}, {n}, {item} — подстановки из игры, их надо сохранить.",
    "Вид: реплика — внизу экрана с именем; мысль — золотым курсивом без имени; летопись — крупно по центру;",
    "заставка — заголовок главы; табличка — имя и роль над головой; субтитр — реплика в бою без остановки игры.",
    "Зелёная колонка «Было раньше» — прежний текст реплики, если она изменилась в последней правке; "
    "«%s» — реплики раньше не было." % NEW_MARK,
    "Порядок жертв выбирает игрок, поэтому сцены дара идут в любом порядке; реплики Ильвы и Сигварда — по номеру дара.",
]
for i, t in enumerate(notes, start=3):
    info.cell(row=i, column=1, value=t).font = Font(name=FONT, size=11)
r0 = len(notes) + 4
info.cell(row=r0, column=1, value="Пример заполненной строки:").font = Font(name=FONT, size=11, bold=True)
ex_head = ["Метка", "Кто говорит", "Текст сейчас", "Новый текст", "Комментарий"]
ex_row = ["gift.sword.plea", "Торстейн", lines["gift.sword.plea"]["text"],
          "Без меча я для них просто мясо, брат. А мне ещё вести своих домой.",
          "Пусть Торстейн хромает, когда выходит."]
for j, (h, v) in enumerate(zip(ex_head, ex_row), start=1):
    hc = info.cell(row=r0 + 1, column=j, value=h)
    hc.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    hc.fill = head_fill
    c = info.cell(row=r0 + 2, column=j, value=v)
    c.font = Font(name=FONT, size=10)
    c.alignment = wrap
    c.border = border
    if j >= 4:
        c.fill = edit_fill
for j, w in enumerate([24, 16, 50, 50, 36], start=1):
    info.column_dimensions[get_column_letter(j)].width = w
info.column_dimensions["A"].width = 24

ws = wb.create_sheet("Диалоги")
headers = ["Метка", "Сцена", "№", "Ситуация (что на экране)", "Условие", "Кто говорит", "Вид", "Текст сейчас",
           "Было раньше", "Новый текст", "Комментарий"]
widths = [26, 26, 5, 52, 26, 16, 12, 60, 46, 60, 36]
EDIT_FROM = headers.index("Новый текст") + 1
BEFORE_COL = headers.index("Было раньше") + 1
for j, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = border
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 30

scene_order = []
counter = {}
changed = added = 0
for i, (label, scene, situation, condition, who, kind) in enumerate(rows, start=2):
    top = scene.split(":")[0]
    if top not in scene_order:
        scene_order.append(top)
    counter[scene] = counter.get(scene, 0) + 1
    fill = PatternFill("solid", fgColor=scene_fills[scene_order.index(top) % len(scene_fills)])
    was = before(label)
    if was == NEW_MARK:
        added += 1
    elif was is not None:
        changed += 1
    values = [label, scene, counter[scene], situation, condition, who, kind, lines[label]["text"], was, None, None]
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=i, column=j, value=v)
        grey = j == 1 or j == BEFORE_COL
        c.font = Font(name=FONT, size=10, bold=(j == 8), color="7F7F7F" if grey else "000000")
        c.alignment = wrap
        c.border = border
        if j >= EDIT_FROM:
            c.fill = edit_fill
        elif j == BEFORE_COL and was is not None:
            c.fill = changed_fill
        else:
            c.fill = fill
ws.freeze_panes = "B2"
ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)

wb.save(out)
print("rows:", len(rows), "changed:", changed, "new:", added, "->", out)
